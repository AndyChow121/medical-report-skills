#!/usr/bin/env python3
"""GRADE Summary of Findings (SoF) 证据摘要表生成器 v1.0

从结构化 JSON 生成 GRADE 标准证据摘要表（Markdown；python-docx 可用时可选 .docx）。
零第三方依赖（除可选 python-docx 降级提示）。

GRADE 质量从高起步，每命中一条降级标准降一级：
    高 -> 中 -> 低 -> 极低
降级维度：偏倚风险 / 不一致性 / 间接性 / 不精确性 / 发表偏倚。

用法：
    python grade_sof.py sof.json [--out sof.md] [--docx sof.docx]
    python grade_sof.py --self-test
"""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

QUALITY_LEVELS = ["高", "中等", "低", "极低"]
EFFECT_TYPES = ["RR", "OR", "HR", "MD", "RD", "SMD"]
BENEFIT_VALUES = ["明确获益", "可能获益", "不确定", "可能有害", "明确有害"]


def read_text(path: Path) -> str:
    for enc in ("utf-8", "gbk", "utf-8-sig"):
        try:
            return path.read_text(encoding=enc)
        except (UnicodeDecodeError, LookupError):
            continue
    return path.read_text(encoding="utf-8", errors="replace")


def suggested_quality(reasons: list[str]) -> str:
    idx = min(len(reasons), len(QUALITY_LEVELS) - 1)
    return QUALITY_LEVELS[idx]


def validate(data: dict) -> list[str]:
    errors = []
    if not isinstance(data.get("outcomes"), list) or not data["outcomes"]:
        errors.append("outcomes 必须为非空数组")
    for i, o in enumerate(data.get("outcomes", [])):
        if not o.get("name"):
            errors.append(f"outcomes[{i}] 缺少 name")
        q = o.get("quality")
        if q is not None and q not in QUALITY_LEVELS:
            errors.append(f"outcomes[{i}] quality 非法: {q!r}（应属 {QUALITY_LEVELS}）")
        et = o.get("effect_type")
        if et is not None and et not in EFFECT_TYPES:
            errors.append(f"outcomes[{i}] effect_type 非法: {et!r}（应属 {EFFECT_TYPES}）")
        bh = o.get("benefit_harm")
        if bh is not None and bh not in BENEFIT_VALUES:
            errors.append(f"outcomes[{i}] benefit_harm 非法: {bh!r}（应属 {BENEFIT_VALUES}）")
    return errors


def render_markdown(data: dict) -> str:
    title = data.get("title", "GRADE 证据摘要表 (Summary of Findings)")
    pop = data.get("population", "")
    iv = data.get("intervention", "")
    cp = data.get("comparison", "")
    lines = [
        f"# {title}",
        "",
        f"**人群(P)**：{pop}　**干预(I)**：{iv}　**对照(C)**：{cp}",
        "",
        "| 结局 | 研究数(受试者) | 效应估计 (95% CI) | 绝对效应 | 证据质量 | 获益/伤害 |",
        "|---|---|---|---|---|---|",
    ]
    for o in data["outcomes"]:
        name = o["name"]
        n_studies = o.get("studies", "—")
        n_part = o.get("participants", "—")
        et = o.get("effect_type", "")
        est = o.get("effect_estimate", "—")
        lo = o.get("ci_low", "")
        hi = o.get("ci_high", "")
        effect = f"{et} {est} ({lo}–{hi})" if et else f"{est} ({lo}–{hi})"
        abs_eff = o.get("absolute_effect", "—")
        quality = o.get("quality", "—")
        reasons = o.get("downgrade_reasons", [])
        if reasons:
            quality += " ↓" + ";".join(reasons)
        bh = o.get("benefit_harm", "—")
        lines.append(f"| {name} | {n_studies} ({n_part}) | {effect} | {abs_eff} | {quality} | {bh} |")
    lines += [
        "",
        "**质量降级说明**：高→中→低→极低，每命中一条降级标准降一级（偏倚风险 / 不一致性 / 间接性 / 不精确性 / 发表偏倚）。",
        "**效应类型**：RR 相对危险度 / OR 比值比 / HR 风险比 / MD 均数差 / RD 风险差 / SMD 标准化均数差。",
        "本表为研究/教育用途，治疗决策需结合临床情境与指南。",
    ]
    return "\n".join(lines)


def consistency_checks(data: dict) -> list[str]:
    """自动建议质量 vs 用户给定质量的一致性提示。"""
    notes = []
    for o in data.get("outcomes", []):
        reasons = o.get("downgrade_reasons", []) or []
        sug = suggested_quality(reasons)
        given = o.get("quality")
        if given and given != sug:
            notes.append(
                f"结局「{o['name']}」给定质量={given}，但按降级理由数({len(reasons)})建议={sug}；请核对。"
            )
    return notes


def to_docx(md_path: Path, docx_path: Path, data: dict) -> str:
    try:
        from docx import Document
    except ImportError:
        return "python-docx 未安装，已仅生成 Markdown（docx 降级跳过）。"
    doc = Document()
    doc.add_heading(data.get("title", "GRADE 证据摘要表"), level=1)
    sub = f"人群(P)：{data.get('population','')}　干预(I)：{data.get('intervention','')}　对照(C)：{data.get('comparison','')}"
    doc.add_paragraph(sub)
    table = doc.add_table(rows=1, cols=6)
    table.style = "Light Grid Accent 1"
    hdr = table.rows[0].cells
    for i, h in enumerate(["结局", "研究数(受试者)", "效应估计 (95% CI)", "绝对效应", "证据质量", "获益/伤害"]):
        hdr[i].text = h
    for o in data["outcomes"]:
        row = table.add_row().cells
        row[0].text = o["name"]
        row[1].text = f"{o.get('studies','—')} ({o.get('participants','—')})"
        et, est, lo, hi = o.get("effect_type", ""), o.get("effect_estimate", "—"), o.get("ci_low", ""), o.get("ci_high", "")
        row[2].text = f"{et} {est} ({lo}–{hi})" if et else f"{est} ({lo}–{hi})"
        row[3].text = o.get("absolute_effect", "—")
        q = o.get("quality", "—")
        reasons = o.get("downgrade_reasons", [])
        if reasons:
            q += " ↓" + ";".join(reasons)
        row[4].text = q
        row[5].text = o.get("benefit_harm", "—")
    doc.save(str(docx_path))
    return f"docx 已生成: {docx_path}"


def run(data: dict, out_md: Path | None, out_docx: Path | None) -> dict:
    errors = validate(data)
    if errors:
        return {"ok": False, "errors": errors}
    md = render_markdown(data)
    notes = consistency_checks(data)
    if out_md:
        out_md.write_text(md + "\n", encoding="utf-8")
    if out_docx:
        note = to_docx(out_md or Path("sof.md"), out_docx, data)
        notes.append(note)
    if out_md is None and out_docx is None:
        print(md)
    return {"ok": True, "markdown": md, "notes": notes}


def _self_test() -> int:
    sample = {
        "title": "他汀 vs 安慰剂 主要心血管事件（SoF 自测）",
        "population": "中危成人", "intervention": "他汀", "comparison": "安慰剂",
        "outcomes": [
            {"name": "全因死亡", "studies": 12, "participants": 45000, "effect_type": "HR",
             "effect_estimate": 0.95, "ci_low": 0.88, "ci_high": 1.03,
             "absolute_effect": "干预 9.1% vs 对照 9.6%（每1000人少5例）",
             "quality": "中等", "downgrade_reasons": ["不精确（CI跨1）"], "benefit_harm": "不确定"},
            {"name": "肌病", "studies": 12, "participants": 45000, "effect_type": "RR",
             "effect_estimate": 1.10, "ci_low": 1.02, "ci_high": 1.19,
             "absolute_effect": "干预 11/1000 vs 对照 10/1000",
             "quality": "高", "downgrade_reasons": [], "benefit_harm": "可能有害"},
        ],
    }
    # 非法枚举：quality 不在 {高,中等,低,极低} -> 应被 validate 捕获
    bad = {
        "outcomes": [
            {"name": "X", "quality": "极高", "benefit_harm": "可能获益"},
        ]
    }
    # 一致性：quality=高 但降级理由 2 条(建议=低) -> 应触发一致性提示
    consistency_sample = {
        "outcomes": [
            {"name": "Y", "quality": "高", "downgrade_reasons": ["偏倚", "不精确"], "benefit_harm": "可能获益"},
        ]
    }
    res_good = run(sample, None, None)
    res_bad = validate(bad)
    ok_good = res_good["ok"]
    ok_bad = len(res_bad) > 0  # bad 应被校验捕获
    consistency = run(consistency_sample, None, None)
    consistency_hit = any("建议" in n for n in consistency.get("notes", []))
    print(f"[self-test] good ok={ok_good} -> {'PASS' if ok_good else 'FAIL'}")
    print(f"[self-test] bad validate captured={len(res_bad)>0} -> {'PASS' if len(res_bad)>0 else 'FAIL'}")
    print(f"[self-test] consistency warn={consistency_hit} -> {'PASS' if consistency_hit else 'FAIL'}")
    return 0 if (ok_good and ok_bad and consistency_hit) else 1


def main() -> int:
    p = argparse.ArgumentParser(description="GRADE SoF 证据摘要表生成器")
    p.add_argument("json_path", nargs="?", help="SoF 数据 JSON")
    p.add_argument("--out", help="Markdown 输出路径")
    p.add_argument("--docx", help="可选 .docx 输出路径（需 python-docx）")
    p.add_argument("--self-test", action="store_true")
    args = p.parse_args()
    if args.self_test:
        return _self_test()
    if not args.json_path:
        p.error("json_path 必填（或用 --self-test）")
    data = json.loads(read_text(Path(args.json_path).expanduser().resolve()))
    res = run(data, Path(args.out) if args.out else None, Path(args.docx) if args.docx else None)
    if not res["ok"]:
        print("校验失败:")
        for e in res["errors"]:
            print(" -", e)
        return 1
    for n in res.get("notes", []):
        print("[note]", n)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
